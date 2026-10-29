# 색 읽기 테스트용 엑셀(가짜 데이터) 만들기 — tests/fixtures/색.xlsx, 무색.xlsx
# 실행: python3 tests/make-color-fixture.py   (openpyxl 필요. SheetJS 무료판은 서식을 못 써서 파이썬으로 만듦)
# 담긴 것: 직접 지정 색(rgb), 테마색+밝기(tint), 옛 색 번호(indexed), 글자색, 제목 줄 색(무시돼야 함),
#          값 없는 칸 색(무시돼야 함), 조건부 서식(못 읽는다고 알려야 함), 칸 순서가 다른 탭의 색
import openpyxl, sys
from openpyxl.styles import PatternFill, Font
from openpyxl.styles.colors import Color
from openpyxl.formatting.rule import CellIsRule
import os
d=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')
H=['번호','상태','금액','메모']
wb=openpyxl.Workbook(); ws=wb.active; ws.title='1월'; ws.append(H)
red=PatternFill('solid', fgColor='FFFF0000')
theme_y=PatternFill('solid', fgColor=Color(theme=7, tint=0.5999938962981048))  # 강조4(FFC000) 밝게 60%
idx_g=PatternFill('solid', fgColor=Color(indexed=42))  # CCFFCC
for i in range(1,7): ws.append([i,'정상',i*100,'메모'+str(i)])
ws['B2'].fill=red; ws['B3'].fill=red; ws['B4'].fill=theme_y; ws['B5'].fill=idx_g
ws['C2'].font=Font(color='FF0070C0'); ws['C3'].font=Font(color=Color(theme=5))  # 강조2 ED7D31
ws['D7'].fill=red; ws['D7'].value=None   # 값 없는 칸 칠
ws['A1'].fill=red  # 제목 줄 색은 무시
ws.conditional_formatting.add('C2:C7', CellIsRule(operator='greaterThan', formula=['300'], fill=red))
ws=wb.create_sheet('2월'); ws.append(['금액','메모','상태','번호'])
for i in range(7,9): ws.append([i*100,'메모'+str(i),'정상',i])
ws['C2'].fill=red
wb.save(f'{d}/색.xlsx')
wb=openpyxl.Workbook(); ws=wb.active; ws.append(H); ws.append([9,'정상',900,'x'])
wb.save(f'{d}/무색.xlsx')
